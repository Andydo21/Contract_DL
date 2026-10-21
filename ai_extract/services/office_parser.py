"""Native Office content with source coordinates; never pretend to render pages."""
import csv
import posixpath
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def extract_docx(path, image_directory, document_key):
    """Read paragraphs, tables and original images in XML document order."""
    from .layout_extractor import save_embedded_image
    ns = dict(NS, w='http://schemas.openxmlformats.org/wordprocessingml/2006/main')
    chunks = []
    with ZipFile(path) as archive:
        document = ET.fromstring(archive.read('word/document.xml'))
        rels = ET.fromstring(archive.read('word/_rels/document.xml.rels')) if 'word/_rels/document.xml.rels' in archive.namelist() else []
        targets = {r.attrib['Id']: posixpath.normpath('word/'+r.attrib['Target']) for r in rels if r.attrib.get('TargetMode') != 'External'}
        styles = {}
        if 'word/styles.xml' in archive.namelist():
            for style in ET.fromstring(archive.read('word/styles.xml')).findall('w:style', ns):
                level = style.find('w:pPr/w:outlineLvl', ns)
                if level is not None:
                    styles[style.attrib.get(f"{{{ns['w']}}}styleId")] = int(level.attrib[f"{{{ns['w']}}}val"])+1
        body = document.find('w:body', ns)
        if body is None:
            return []
        def add(text, kind, location, **metadata):
            chunks.append(dict(chunk_id=len(chunks)+1, text=text, layout_type=kind, page_number=1,
                               bbox=[], source_location=location, location_kind='document_block',
                               has_image=bool(metadata.get('image_path')), extractor_version='native_office_v2', **metadata))
        for index, element in enumerate(body, 1):
            location = f'Block {index}'
            if element.tag == f"{{{ns['w']}}}p":
                text = ''.join(t.text or '' for t in element.findall('.//w:t', ns))
                style = element.find('w:pPr/w:pStyle', ns)
                style_id = style.attrib.get(f"{{{ns['w']}}}val", '') if style is not None else ''
                level = styles.get(style_id)
                if level is None:
                    import re
                    match = re.fullmatch(r'Heading([1-9])', style_id, re.I)
                    level = int(match.group(1)) if match else None
                if text.strip():
                    add(text, 'title' if level else 'paragraph', location,
                        **(dict(heading_level=level, heading_source='native_style') if level else {}))
            elif element.tag == f"{{{ns['w']}}}tbl":
                rows = [[''.join(t.text or '' for t in cell.findall('.//w:t', ns)) for cell in row.findall('w:tc', ns)] for row in element.findall('w:tr', ns)]
                headers = rows[0] if rows else []
                for row_index, row in enumerate(rows, 1):
                    add(' | '.join(row), 'table', f'{location}, row {row_index}', table_headers=headers,
                        table_values=row, table_row=row_index, table_id=f'word-table:{index}',
                        table_header_origin='source_first_row_unverified')
            for drawing in element.findall('.//a:blip', ns):
                reference = drawing.attrib.get(f"{{{ns['r']}}}embed")
                target = targets.get(reference)
                if not target or target not in archive.namelist():
                    continue
                blob = archive.read(target)
                try:
                    image = save_embedded_image(blob, image_directory, document_key)
                    add('', 'figure', location, image_path='extracted_images/'+image.name, visual_status='original_embedded_image')
                except (ValueError, OSError):
                    add('', 'figure', location, visual_status='unsupported_embedded_image')
    return chunks


def extract_pptx(path, image_directory=None, document_key=None):
    chunks = []
    with ZipFile(path) as archive:
        presentation = ET.fromstring(archive.read('ppt/presentation.xml'))
        rels = ET.fromstring(archive.read('ppt/_rels/presentation.xml.rels'))
        targets = {r.attrib['Id']: r.attrib['Target'] for r in rels
                   if r.attrib.get('TargetMode') != 'External'}
        for number, slide_id in enumerate(presentation.findall('p:sldIdLst/p:sldId', NS), 1):
            target = targets[slide_id.attrib[f"{{{NS['r']}}}id"]]
            name = target.lstrip('/') if target.startswith('/') else posixpath.normpath('ppt/' + target)
            slide = ET.fromstring(archive.read(name))
            rel_name = posixpath.dirname(name) + '/_rels/' + posixpath.basename(name) + '.rels'
            slide_rels = ET.fromstring(archive.read(rel_name)) if rel_name in archive.namelist() else []
            images = {r.attrib['Id']: posixpath.normpath(posixpath.dirname(name)+'/'+r.attrib['Target'])
                      for r in slide_rels if r.attrib.get('TargetMode') != 'External'}
            tree = slide.find('p:cSld/p:spTree', NS)
            if tree is None:
                continue
            for shape in tree.iter():
                if shape.tag == f"{{{NS['p']}}}sp":
                    paragraphs = [''.join(t.text or '' for t in p.findall('.//a:t', NS))
                                  for p in shape.findall('p:txBody/a:p', NS)]
                    text = '\n'.join(p for p in paragraphs if p.strip())
                    kind = 'paragraph'
                elif shape.tag == f"{{{NS['p']}}}graphicFrame":
                    table = shape.find('.//a:tbl', NS)
                    if table is None:
                        continue
                    rows = [[' '.join(t.text or '' for t in cell.findall('.//a:t', NS))
                             for cell in row.findall('a:tc', NS)] for row in table.findall('a:tr', NS)]
                    text = '\n'.join(' | '.join(row) for row in rows)
                    kind = 'table'
                elif shape.tag == f"{{{NS['p']}}}pic":
                    blip = shape.find('.//a:blip', NS)
                    target = images.get(blip.attrib.get(f"{{{NS['r']}}}embed")) if blip is not None else None
                    if image_directory and target in archive.namelist():
                        chunks.append(embedded_chunk(archive.read(target), image_directory, document_key,
                                      len(chunks)+1, number, f'Slide {number}', 'slide'))
                    continue
                else:
                    continue
                if text.strip():
                    chunks.append(dict(chunk_id=len(chunks) + 1, layout_type=kind,
                        text=text, page_number=number, source_location=f'Slide {number}',
                        location_kind='slide', bbox=[], has_image=False, image_url=None,
                        extractor_version='native_office_v1'))
    return chunks


def extract_spreadsheet(path, image_directory=None, document_key=None):
    suffix = Path(path).suffix.lower()
    if suffix in {'.csv', '.tsv'}:
        with open(path, encoding='utf-8-sig', newline='') as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel_tab if suffix == '.tsv' else csv.excel
            rows = list(csv.reader(handle, dialect))
        return row_chunks(Path(path).name, rows, 1)
    if suffix != '.xlsx':
        raise ValueError('Hãy chuyển bảng tính sang XLSX, CSV hoặc TSV trước khi trích xuất.')
    from openpyxl import load_workbook
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    chunks = []
    try:
        for index, sheet in enumerate(values, 1):
            merged = [str(r) for r in sheet.merged_cells.ranges]
            anchors = {}
            for region in sheet.merged_cells.ranges:
                anchor = sheet.cell(region.min_row, region.min_col).value
                for row in range(region.min_row, region.max_row + 1):
                    for col in range(region.min_col, region.max_col + 1):
                        anchors[(row, col)] = anchor
            rows = []
            formula_map = {}
            for row in sheet.iter_rows():
                result = []
                for cell in row:
                    value = cell.value if cell.value is not None else anchors.get((cell.row, cell.column))
                    formula = formulas[sheet.title].cell(cell.row, cell.column)
                    if formula.data_type == 'f':
                        formula_map[formula.coordinate] = formula.value
                        if value is None:
                            value = f'[Formula {formula.value}; chưa có giá trị cache]'
                    result.append('' if value is None else str(value))
                rows.append(result)
            sheet_chunks = row_chunks(sheet.title, rows, index)
            for chunk in sheet_chunks:
                chunk.update(merged_ranges=merged, formulas={k: v for k, v in formula_map.items()
                             if k in chunk['cell_values']})
                chunk['chunk_id'] = len(chunks) + 1
                chunks.append(chunk)
            if image_directory:
                for embedded in sheet._images:
                    anchor = getattr(embedded.anchor, '_from', None)
                    cell = f'{column_name(anchor.col+1)}{anchor.row+1}' if anchor is not None else 'unknown anchor'
                    chunks.append(embedded_chunk(embedded._data(), image_directory, document_key,
                                  len(chunks)+1, index, f'Sheet {sheet.title}, image at {cell}', 'sheet_image'))
    finally:
        values.close()
        formulas.close()
    return chunks


def column_name(number):
    name = ''
    while number:
        number, mod = divmod(number - 1, 26)
        name = chr(65 + mod) + name
    return name


def row_chunks(sheet_name, rows, sheet_number):
    first = next((i for i, row in enumerate(rows) if any(str(c).strip() for c in row)), None)
    if first is None:
        return []
    headers = [str(c) for c in rows[first]]
    result = []
    for index in range(first, len(rows)):
        row = rows[index]
        if not any(str(c).strip() for c in row):
            continue
        cells = {f'{column_name(col)}{index + 1}': str(value) for col, value in enumerate(row, 1)}
        text = '; '.join(f'{headers[col - 1] if col <= len(headers) and headers[col - 1] else column_name(col)}: {value}'
                         for col, value in enumerate(row, 1) if str(value).strip())
        result.append(dict(chunk_id=len(result) + 1, layout_type='table',
            text=f'[Sheet: {sheet_name} | Row: {index + 1}]\n{text}',
            sheet_name=sheet_name, table_id=f'sheet:{sheet_number}', table_headers=headers,
            table_row=index + 1, table_values=[str(value) for value in row],
            table_header_origin='source_first_row_unverified', cell_values=cells, bbox=[], has_image=False, image_url=None,
            page_number=sheet_number, source_location=f'Sheet {sheet_name}, row {index + 1}',
            location_kind='sheet_row', extractor_version='native_office_v1'))
    return result


def embedded_chunk(blob, directory, document_key, chunk_id, page_number, location, location_kind):
    from .layout_extractor import save_embedded_image
    chunk = dict(chunk_id=chunk_id, text='', layout_type='figure', page_number=page_number,
                 bbox=[], source_location=location, location_kind=location_kind, has_image=False,
                 extractor_version='native_office_v2')
    try:
        path = save_embedded_image(blob, directory, document_key)
        chunk.update(image_path='extracted_images/'+path.name, has_image=True,
                     visual_status='original_embedded_image')
    except (ValueError, OSError):
        chunk['visual_status'] = 'unsupported_embedded_image'
    return chunk
